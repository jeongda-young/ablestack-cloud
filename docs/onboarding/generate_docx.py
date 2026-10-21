#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the editable Mold onboarding handbook from its Markdown source.

Run with the bundled Codex Python runtime (python-docx and Pillow). Images are
read from the source directory; this builder never regenerates diagrams.
"""

import argparse
import hashlib
import html
import json
import re
import struct
import textwrap
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Mm, Pt, RGBColor
from PIL import Image
from lxml import etree


FONT_NAME = "NanumSquare"
BLACK = "000000"
INK = "202020"
NAVY = "203A56"
RULE = "D9D9D9"
PALE = "F3F6F8"
BODY_SIZE = 10.5
HERE = Path(__file__).resolve().parent


def embed_document_fonts(docx_path, font_dir=None):
    """Embed complete, unmodified NanumSquare faces using OOXML obfuscation.

    The distributed font files permit Preview & Print embedding (fsType=4).
    Do not change those permissions: a recipient without installed fonts may
    need to install NanumSquare separately to edit in Word.
    """
    docx_path = Path(docx_path)
    font_dir = Path(font_dir) if font_dir else HERE / "fonts"
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pr = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct = "http://schemas.openxmlformats.org/package/2006/content-types"
    with zipfile.ZipFile(docx_path) as archive:
        infos = archive.infolist()
        parts = {info.filename: archive.read(info) for info in infos}
    table = etree.fromstring(parts["word/fontTable.xml"])
    font = table.find(f"{{{w}}}font[@{{{w}}}name='{FONT_NAME}']")
    if font is None:
        font = etree.SubElement(table, f"{{{w}}}font", {f"{{{w}}}name": FONT_NAME})
        etree.SubElement(font, f"{{{w}}}family", {f"{{{w}}}val": "swiss"})
        etree.SubElement(font, f"{{{w}}}pitch", {f"{{{w}}}val": "variable"})
    rel_path = "word/_rels/fontTable.xml.rels"
    rels = (etree.fromstring(parts[rel_path]) if rel_path in parts else
            etree.Element(f"{{{pr}}}Relationships", nsmap={None: pr}))
    for face, filename in (("Regular", "NanumSquareR.ttf"),
                           ("Bold", "NanumSquareB.ttf")):
        raw = (font_dir / filename).read_bytes()
        if raw[:4] != b"\x00\x01\x00\x00":
            raise ValueError(f"Expected TrueType font: {filename}")
        records = [struct.unpack_from(">4sIII", raw, 12 + i * 16)
                   for i in range(struct.unpack_from(">H", raw, 4)[0])]
        os2 = next(record[2] for record in records if record[0] == b"OS/2")
        permissions = struct.unpack_from(">H", raw, os2 + 8)[0]
        if permissions & 0x202 or (permissions & 15) not in (0, 4, 8):
            raise ValueError(f"Font does not permit outline embedding: {filename}")
        key = uuid.uuid5(uuid.NAMESPACE_URL, hashlib.sha256(raw).hexdigest())
        mask = key.bytes[::-1]
        encoded = bytearray(raw)
        for index in range(32):
            encoded[index] ^= mask[index % 16]
        target = "fonts/" + Path(filename).stem + ".odttf"
        parts["word/" + target] = bytes(encoded)
        relation = next((node for node in rels
                         if node.get("Target") == target), None)
        if relation is None:
            used_ids = {node.get("Id") for node in rels}
            number = 1
            while f"rId{number}" in used_ids:
                number += 1
            relation = etree.SubElement(rels, f"{{{pr}}}Relationship", {
                "Id": f"rId{number}", "Type": r + "/font", "Target": target})
        embedded = font.find(f"{{{w}}}embed{face}")
        if embedded is None:
            embedded = etree.SubElement(font, f"{{{w}}}embed{face}")
        embedded.set(f"{{{r}}}id", relation.get("Id"))
        embedded.set(f"{{{w}}}fontKey", "{" + str(key).upper() + "}")
        embedded.set(f"{{{w}}}subsetted", "false")
    settings = etree.fromstring(parts["word/settings.xml"])
    # These settings precede proofState/defaultTabStop in CT_Settings order.
    preceding = {"writeProtection", "view", "zoom", "removePersonalInformation",
                 "removeDateAndTime", "doNotDisplayPageBoundaries",
                 "displayBackgroundShape", "printPostScriptOverText",
                 "printFractionalCharacterWidth", "printFormsData"}
    for name in ("embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts"):
        for node in settings.findall(f"{{{w}}}{name}"):
            settings.remove(node)
    position = 0
    while position < len(settings) and etree.QName(settings[position]).localname in preceding:
        position += 1
    for offset, (name, value) in enumerate((('embedTrueTypeFonts', 'true'),
                                           ('embedSystemFonts', 'false'),
                                           ('saveSubsetFonts', 'false'))):
        settings.insert(position + offset, etree.Element(
            f"{{{w}}}{name}", {f"{{{w}}}val": value}))
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    for extension, mime in (("odttf", "application/vnd.openxmlformats-officedocument.obfuscatedFont"),
                            ("txt", "text/plain")):
        if content_types.find(f"{{{ct}}}Default[@Extension='{extension}']") is None:
            etree.SubElement(content_types, f"{{{ct}}}Default", {
                "Extension": extension, "ContentType": mime})
    parts["word/fonts/LICENSE.txt"] = (font_dir / "LICENSE.txt").read_bytes()
    for path, root in (("word/fontTable.xml", table), (rel_path, rels),
                       ("word/settings.xml", settings),
                       ("[Content_Types].xml", content_types)):
        parts[path] = etree.tostring(root, xml_declaration=True,
                                    encoding="UTF-8", standalone=True)
    temp_path = docx_path.with_suffix(".font-embedding.tmp")
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for info in infos:
                archive.writestr(info, parts.pop(info.filename))
            for path, content in parts.items():
                archive.writestr(path, content)
        temp_path.replace(docx_path)
    finally:
        temp_path.unlink(missing_ok=True)


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    data: object = None
    anchor: str = ""


def ensure_child(parent, tag):
    element = parent.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        parent.append(element)
    return element


def set_rpr_font(rpr, size=None, color=None, bold=None, italic=None):
    fonts = ensure_child(rpr, "w:rFonts")
    for name in list(fonts.attrib):
        del fonts.attrib[name]
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), FONT_NAME)
    if size is not None:
        for tag in ("w:sz", "w:szCs"):
            ensure_child(rpr, tag).set(qn("w:val"), str(round(size * 2)))
    if color:
        node = ensure_child(rpr, "w:color")
        node.attrib.clear()
        node.set(qn("w:val"), color)
    for tag, value in (("w:b", bold), ("w:bCs", bold),
                       ("w:i", italic), ("w:iCs", italic)):
        if value is not None:
            ensure_child(rpr, tag).set(qn("w:val"), "1" if value else "0")


def apply_font(run, size=BODY_SIZE, color=INK, bold=False, italic=False):
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    set_rpr_font(run._r.get_or_add_rPr(), size, color, bold, italic)


def clean_heading(text):
    """Use plain heading text while retaining source body text verbatim."""
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = text.replace("`", "").replace("**", "")
    # Keep section numbers distinct using the requested 1-1 / A-1 display.
    # Extract them before removing display-only punctuation from the title.
    number = re.match(r"^((?:\d+|[A-Z])(?:\.\d+)+)\s+(.+)$", text)
    prefix = ""
    if number:
        prefix = number.group(1).replace(".", "-") + " "
        text = number.group(2)
    # Unicode letters/numbers retain Korean, English, and technical names.
    chars = [ch if unicodedata.category(ch)[0] in "LN" or ch.isspace() else " "
             for ch in text]
    return prefix + re.sub(r"\s+", " ", "".join(chars)).strip()


def display_text(text):
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    return html.unescape(text.replace("**", "").replace("`", ""))


def is_content_lead(text):
    """Keep a standalone subheading or lead attached to its explanation."""
    return bool(re.fullmatch(r"\*\*[^*\n]+\*\*", text.strip())) or display_text(text).rstrip().endswith((
        ":", "：", "완료 기준은 다음과 같다.", "완료 기준은 다음과 같습니다."
    ))


def heading_role(block, seen_title):
    if block.level == 1 and not seen_title:
        return "Title"
    if block.level == 1:
        return "Heading 1"
    if block.level == 2 and not seen_title:
        return "Subtitle"
    if block.level == 2:
        return "Handbook Frontmatter"
    if block.level == 3 and re.match(r"^(?:\d+|[A-Z])\.\d+\b", block.text):
        return "Heading 2"
    return "Heading 3"


def split_table_row(line):
    """Split pipes outside code spans and preserve escaped literal pipes."""
    text = line.strip().strip("|")
    cells, current = [], []
    code = False
    i = 0
    while i < len(text):
        char = text[i]
        if char == "\\" and i + 1 < len(text) and text[i + 1] == "|":
            current.append("|")
            i += 2
            continue
        if char == "`":
            code = not code
        if char == "|" and not code:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
        i += 1
    cells.append("".join(current).strip())
    return cells


def parse_markdown(text):
    lines = text.splitlines()
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or re.fullmatch(r"(?:-{3,}|\*{3,}|_{3,})", stripped):
            i += 1
            continue
        fence = re.match(r"^\s*(`{3,}|~{3,})(.*)$", line)
        if fence:
            delimiter = fence.group(1)
            collected = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(delimiter):
                collected.append(lines[i])
                i += 1
            blocks.append(Block("code", textwrap.dedent("\n".join(collected))))
            i += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            title = heading.group(2)
            if "전체 목차" in title or "Table of Contents" in title:
                blocks.append(Block("toc"))
                i += 1
                while i < len(lines):
                    current = lines[i].strip()
                    if current.startswith("#") or re.fullmatch(r"-{3,}", current):
                        break
                    i += 1
                continue
            blocks.append(Block("heading", title, len(heading.group(1))))
            i += 1
            continue
        image = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", stripped)
        if image:
            blocks.append(Block("image", image.group(1), data=image.group(2)))
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while i < len(lines):
                current = lines[i].strip()
                if not (current.startswith("|") and current.endswith("|")):
                    break
                cells = split_table_row(current)
                if not all(re.fullmatch(r":?-{2,}:?", cell.replace(" ", ""))
                           for cell in cells):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(Block("table", data=rows))
            continue
        if stripped.startswith(">"):
            notes = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                content = re.sub(r"^\s*>\s?", "", lines[i]).strip()
                content = re.sub(r"\[!(?:NOTE|IMPORTANT|CAUTION|WARNING|TIP)\]", "", content).strip()
                if content:
                    notes.append(content)
                i += 1
            blocks.append(Block("note", data=notes))
            continue
        bullet = re.match(r"^(\s*)[-*+]\s+(.*)$", line)
        if bullet:
            blocks.append(Block("bullet", bullet.group(2),
                                level=min(2, len(bullet.group(1)) // 2)))
            i += 1
            continue
        number = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if number:
            blocks.append(Block("number", number.group(3),
                                level=min(2, len(number.group(1)) // 2),
                                data=number.group(2)))
            i += 1
            continue
        # One Markdown source line remains one ordinary Word paragraph.
        blocks.append(Block("paragraph", stripped))
        i += 1
    return blocks


def append_link(paragraph, text, target, size=BODY_SIZE, color=NAVY, bold=False):
    node = OxmlElement("w:hyperlink")
    node.set(qn("w:history"), "1")
    if target.startswith("#"):
        node.set(qn("w:anchor"), target[1:])
    else:
        rel = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK,
                                       is_external=True)
        node.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    set_rpr_font(rpr, size, color, bold, False)
    run.append(rpr)
    content = OxmlElement("w:t")
    content.text = text
    content.set(qn("xml:space"), "preserve")
    run.append(content)
    node.append(run)
    paragraph._p.append(node)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^]]+\]\([^)]+\)|<br\s*/?>)")


def add_inline(paragraph, text, size=BODY_SIZE, color=INK, bold=False):
    """Preserve emphasis and links without introducing a second font."""
    for part in INLINE.split(text):
        if not part:
            continue
        if re.fullmatch(r"<br\s*/?>", part):
            run = paragraph.add_run()
            apply_font(run, size, color, bold)
            run.add_break()
            continue
        if part.startswith("**") and part.endswith("**"):
            add_inline(paragraph, part[2:-2], size, color, True)
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(html.unescape(part[1:-1]))
            apply_font(run, size, color, bold)
            continue
        link = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", part)
        if link:
            append_link(paragraph, link.group(1), link.group(2), size,
                        color if color == "FFFFFF" else NAVY, bold)
            continue
        run = paragraph.add_run(html.unescape(part).replace("  \n", "\n"))
        apply_font(run, size, color, bold)


def add_bookmark(paragraph, identifier, name):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(1 if paragraph._p.find(qn("w:pPr")) is not None else 0, start)
    paragraph._p.append(end)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(.75)
    section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.3)

    # Apply the family to every built-in style as well as explicit runs.
    for style in doc.styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER,
                          WD_STYLE_TYPE.TABLE):
            set_rpr_font(ensure_child(style.element, "w:rPr"))
    normal = doc.styles["Normal"]
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    settings = {
        "Title": (25, 14, 12),
        "Subtitle": (12, 0, 14),
        "Heading 1": (18, 0, 12),
        "Heading 2": (13, 16, 8),
        "Heading 3": (11, 12, 6),
        "Caption": (8.5, 2, 8),
    }
    for name, (size, before, after) in settings.items():
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.bold = name != "Subtitle" and name != "Caption"
        set_rpr_font(ensure_child(style.element, "w:rPr"), size, BLACK,
                     style.font.bold, False)
        pf = style.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.25
        pf.keep_with_next = name != "Caption"
        pf.keep_together = True
        # Remove any inherited paragraph borders and theme coloring.
        for borders in style.element.findall(".//" + qn("w:pBdr")):
            borders.getparent().remove(borders)
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.line_spacing = 1.4
        doc.styles[name].paragraph_format.space_after = Pt(5)
    for name in ("Handbook Code", "Handbook Frontmatter", "TOC 1", "TOC 2"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        doc.styles[name].base_style = normal
    code = doc.styles["Handbook Code"]
    code.font.size = Pt(9)
    code.paragraph_format.left_indent = Inches(.12)
    code.paragraph_format.right_indent = Inches(.12)
    code.paragraph_format.line_spacing = 1.25
    code.paragraph_format.space_before = Pt(8)
    code.paragraph_format.space_after = Pt(10)
    code.paragraph_format.widow_control = True
    # Markdown fenced code remains editable text, grouped in a distinct code
    # area. This styling applies only to examples, never explanatory notes.
    code_props = ensure_child(code.element, "w:pPr")
    shading = ensure_child(code_props, "w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F3F5F7")
    borders = ensure_child(code_props, "w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = ensure_child(borders, "w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"),
                           ("space", "6"), ("color", RULE)):
            border.set(qn("w:" + key), value)
    # Adjacent template sections share their enclosing code border.
    ensure_child(borders, "w:between").set(qn("w:val"), "nil")
    for node in (borders, shading):
        code_props.remove(node)
        code_props.insert_element_before(node, "w:tabs", "w:spacing", "w:ind",
                                         "w:contextualSpacing", "w:jc")
    frontmatter = doc.styles["Handbook Frontmatter"]
    frontmatter.font.size = Pt(12)
    frontmatter.font.bold = True
    frontmatter.font.color.rgb = RGBColor.from_string(BLACK)
    frontmatter.paragraph_format.space_before = Pt(12)
    frontmatter.paragraph_format.space_after = Pt(7)
    frontmatter.paragraph_format.keep_with_next = True
    for name in ("TOC 1", "TOC 2"):
        style = doc.styles[name]
        style.font.size = Pt(10.5 if name == "TOC 1" else 9.5)
        style.font.bold = name == "TOC 1"
        style.font.color.rgb = RGBColor.from_string(NAVY)
        pf = style.paragraph_format
        pf.line_spacing = 1.3
        pf.space_after = Pt(4)
        pf.space_before = Pt(8 if name == "TOC 1" else 0)
        pf.keep_together = True
        pf.keep_with_next = name == "TOC 1"
        pf.left_indent = Inches(0 if name == "TOC 1" else .18)
        pf.right_indent = Inches(.3)
        pf.tab_stops.add_tab_stop(Mm(171.9), WD_TAB_ALIGNMENT.RIGHT,
                                 WD_TAB_LEADER.DOTS)

    defaults = ensure_child(doc.styles.element, "w:docDefaults")
    rdefaults = ensure_child(defaults, "w:rPrDefault")
    set_rpr_font(ensure_child(rdefaults, "w:rPr"), BODY_SIZE, INK)
    # The cached native TOC is usable immediately and can be refreshed in Word.
    ensure_child(doc.settings.element, "w:updateFields").set(qn("w:val"), "true")
    section.different_first_page_header_footer = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    r = footer.add_run("Mold 신입 개발자 교육 안내서   |   ")
    apply_font(r, 8, BLACK)
    for field_type in ("begin", "separate", "end"):
        run = footer.add_run()
        apply_font(run, 8, BLACK)
        if field_type == "begin":
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), "begin")
            run._r.append(field)
            instruction_run = footer.add_run()
            apply_font(instruction_run, 8, BLACK)
            instruction = OxmlElement("w:instrText")
            instruction.set(qn("xml:space"), "preserve")
            instruction.text = " PAGE "
            instruction_run._r.append(instruction)
        else:
            field = OxmlElement("w:fldChar")
            field.set(qn("w:fldCharType"), field_type)
            run._r.append(field)
            if field_type == "separate":
                cached = footer.add_run("1")
                apply_font(cached, 8, BLACK)


def add_table(doc, rows):
    count = max(len(row) for row in rows)
    rows = [row + [""] * (count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    available = 210 / 25.4 - 1.5
    # Intentionally reserve narrow fields for step numbers/status and give
    # prose columns space based on their actual header/body lengths.
    scores = []
    for col in range(count):
        lengths = [sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1
                       for ch in display_text(row[col])) for row in rows]
        mean = sum(lengths) / len(lengths)
        score = max(8, min(44, mean * .6 + max(lengths) * .25))
        if max(lengths) <= 14:
            score = max(7, min(14, max(lengths)))
        scores.append(score)
    if rows[0] == ["메뉴", "무엇을 관리하는가", "처음 확인할 내용"]:
        scores[0] = max(scores[0], 22)
    # Reserve readable widths for short label fields (including CCVM/KVM),
    # then redistribute the remaining width across longer fields by weight.
    minimum = min(.9, available / count)
    widths = [0.0] * count
    unfixed = list(range(count))
    remaining = available
    while unfixed:
        weight = sum(scores[column] for column in unfixed)
        narrow = [column for column in unfixed
                  if remaining * scores[column] / weight < minimum]
        if not narrow:
            for column in unfixed:
                widths[column] = remaining * scores[column] / weight
            break
        for column in narrow:
            widths[column] = minimum
            remaining -= minimum
            unfixed.remove(column)
    for col, width in enumerate(widths):
        table.columns[col].width = Inches(width)
    props = table._tbl.tblPr
    borders = ensure_child(props, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = ensure_child(borders, "w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"),
                           ("color", RULE), ("space", "0")):
            node.set(qn("w:" + key), value)
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        rowprops = row._tr.get_or_add_trPr()
        ensure_child(rowprops, "w:cantSplit")
        if row_index == 0:
            ensure_child(rowprops, "w:tblHeader")
        for column, value in enumerate(values):
            cell = row.cells[column]
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cellprops = cell._tc.get_or_add_tcPr()
            ensure_child(cellprops, "w:shd").set(qn("w:fill"),
                        NAVY if row_index == 0 else PALE if row_index % 2 else "FFFFFF")
            margins = ensure_child(cellprops, "w:tcMar")
            for side, value_twips in (("top", 100), ("bottom", 100),
                                     ("left", 100), ("right", 100)):
                node = ensure_child(margins, "w:" + side)
                node.set(qn("w:w"), str(value_twips))
                node.set(qn("w:type"), "dxa")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            # Keep small value columns centered; explanatory columns left.
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if scores[column] <= 14 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, 9.5 if row_index == 0 else 9.2,
                       "FFFFFF" if row_index == 0 else INK,
                       row_index == 0)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(4)
    after.paragraph_format.line_spacing = .5
    apply_font(after.add_run(""), 4)


def add_image(doc, block, source_dir):
    image_path = source_dir / str(block.data)
    if not image_path.is_file():
        raise FileNotFoundError(f"Markdown image not found: {image_path}")
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    width = min(6.4, 5.5 * pixel_width / pixel_height)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = bool(block.text)
    run = p.add_run()
    apply_font(run)
    run.add_picture(str(image_path), width=Inches(width))
    # Preserve the caption as accessible drawing alternative text.
    for node in run._r.findall(".//" + qn("wp:docPr")):
        node.set("descr", block.text)
    if block.text:
        caption = doc.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(caption, block.text, 8.5, BLACK)


def field_run(parent, kind=None, instruction=None, text=None, size=9.5, bold=False):
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    set_rpr_font(rpr, size, NAVY, bold, False)
    run.append(rpr)
    if kind:
        field = OxmlElement("w:fldChar")
        field.set(qn("w:fldCharType"), kind)
        run.append(field)
    elif instruction is not None:
        field = OxmlElement("w:instrText")
        field.set(qn("xml:space"), "preserve")
        field.text = instruction
        run.append(field)
    elif text is not None:
        field = OxmlElement("w:t")
        field.set(qn("xml:space"), "preserve")
        field.text = text
        run.append(field)
    else:
        run.append(OxmlElement("w:tab"))
    parent.append(run)


def render_toc(doc, entries, page_map):
    heading = doc.add_paragraph(style="Handbook Frontmatter")
    heading.paragraph_format.page_break_before = True
    add_inline(heading, "전체 목차", 18, BLACK, True)
    p = doc.add_paragraph()
    add_inline(p, "제목 또는 오른쪽 페이지 번호를 선택하면 해당 절로 이동합니다.", 9.5, BLACK)
    p.paragraph_format.space_after = Pt(12)
    for index, (block, style) in enumerate(entries):
        major = style == "Heading 1"
        size = 10.5 if major else 9.5
        p = doc.add_paragraph(style="TOC 1" if major else "TOC 2")
        # Start coherent chapter groups on separate contents pages.
        if major and (re.match(r"^CHAPTER [358]\b", block.text) or
                      block.text.startswith("부록 A")):
            p.paragraph_format.page_break_before = True
        if index == 0:
            field_run(p._p, kind="begin")
            field_run(p._p, instruction=' TOC \\o "1-2" \\h \\z ')
            field_run(p._p, kind="separate")
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:history"), "1")
        link.set(qn("w:anchor"), block.anchor)
        field_run(link, text=clean_heading(block.text), size=size, bold=major)
        field_run(link, size=size, bold=major)
        field_run(link, kind="begin", size=size, bold=major)
        field_run(link, instruction=f' PAGEREF {block.anchor} \\h ', size=size, bold=major)
        field_run(link, kind="separate", size=size, bold=major)
        field_run(link, text=str(page_map.get(block.anchor, 1)), size=size, bold=major)
        field_run(link, kind="end", size=size, bold=major)
        p._p.append(link)
        if index == len(entries) - 1:
            field_run(p._p, kind="end")


def enforce_font(doc):
    """Set explicit family on all runs, including fields/hyperlinks/drawings."""
    parts = [doc.part]
    parts.extend(rel.target_part for rel in doc.part.rels.values()
                 if not rel.is_external and rel.reltype in
                 (RELATIONSHIP_TYPE.HEADER, RELATIONSHIP_TYPE.FOOTER))
    for part in parts:
        for run in part.element.iter(qn("w:r")):
            set_rpr_font(ensure_child(run, "w:rPr"))
    for style in doc.styles:
        set_rpr_font(ensure_child(style.element, "w:rPr"))
    # Built-in table styles include nested conditional run formatting. Those
    # descendants can override the style's top-level family and must also be
    # normalized, as must any defaults supplied by the base Word template.
    for fonts in doc.styles.element.iter(qn("w:rFonts")):
        fonts.attrib.clear()
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + name), FONT_NAME)


def add_cover(doc, logo_path):
    """Add one editable cover page using the official company logo."""
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo.paragraph_format.space_before = Pt(40)
    logo.paragraph_format.space_after = Pt(110)
    logo.paragraph_format.keep_together = True
    picture = logo.add_run().add_picture(str(logo_path), width=Mm(65))
    picture._inline.docPr.set("descr", "에이블클라우드 컬러 회사 로고")
    # Keep the official GitHub logo bytes intact; crop outer padding in Word.
    # Retain the previous 65 mm picture frame so cover spacing stays stable.
    picture.height = Mm(65 * 427 / 2400)
    crop = OxmlElement("a:srcRect")
    for name, value in {"l": "2174", "t": "42826", "r": "2609", "b": "40217"}.items():
        crop.set(name, value)
    picture._inline.xpath(".//pic:blipFill")[0].insert(1, crop)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.line_spacing = 1.16
    run = title.add_run("Mold")
    apply_font(run, 44, BLACK, True)
    run.add_break()
    apply_font(title.add_run("신입 개발자 교육 안내서"), 28, BLACK, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    subtitle.paragraph_format.space_after = Pt(0)
    add_inline(subtitle, "개발 환경 준비부터 첫 개발 실습까지", 13, INK)

    publisher = doc.add_paragraph()
    publisher.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    publisher.paragraph_format.space_before = Pt(240)
    publisher.paragraph_format.space_after = Pt(8)
    add_inline(publisher, "에이블클라우드", 12, BLACK, True)
    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    team.paragraph_format.space_after = Pt(12)
    add_inline(team, "연구개발", 10.5, INK)
    edition = doc.add_paragraph()
    edition.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    edition.paragraph_format.space_after = Pt(0)
    add_inline(edition, "2026년 10월  |  v1.0", 10, INK)


def create_styled_document(md_path, docx_path, page_map=None):
    md_path, docx_path = Path(md_path), Path(docx_path)
    blocks = parse_markdown(md_path.read_text(encoding="utf-8"))
    doc = Document()
    configure_document(doc)
    add_cover(doc, md_path.parent / "ablecloud-logo-color.png")
    entries = []
    seen_title = False
    for index, block in enumerate(blocks):
        if block.kind != "heading":
            continue
        style = heading_role(block, seen_title)
        # The first H2 is the subtitle immediately after the document title.
        if index == 1 and block.level == 2:
            style = "Subtitle"
        block.data = style
        if style == "Title":
            seen_title = True
        block.anchor = f"section_{index}"
        if style == "Heading 1" or (style == "Heading 2" and
                re.match(r"^(?:\d+|[A-Z])\.\d+\b", block.text)):
            entries.append((block, style))
    in_menu_chapter = False
    for identifier, block in enumerate(blocks, 1):
        if block.kind == "heading":
            if block.data == "Heading 1":
                in_menu_chapter = bool(re.match(r"^CHAPTER 2\b", block.text))
            p = doc.add_paragraph(style=block.data)
            size = doc.styles[block.data].font.size.pt
            add_inline(p, clean_heading(block.text), size, BLACK,
                       block.data != "Subtitle")
            add_bookmark(p, identifier, block.anchor)
            if block.data == "Title" or re.match(r"^(?:9\.[24567]|A\.2|E\.3)\b", block.text):
                p.paragraph_format.page_break_before = True
        elif block.kind == "toc":
            render_toc(doc, entries, page_map or {})
        elif block.kind == "table":
            add_table(doc, block.data)
            if in_menu_chapter and len(block.data) <= 4:
                for row in doc.tables[-1].rows[:-1]:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.keep_with_next = True
        elif block.kind == "image":
            add_image(doc, block, md_path.parent)
        elif block.kind == "code":
            # One fenced source example is one visible, editable code area.
            p = doc.add_paragraph(style="Handbook Code")
            p.paragraph_format.keep_together = len(block.text.splitlines()) <= 24
            run = p.add_run(block.text)
            apply_font(run, 9)
        elif block.kind == "note":
            for index, text in enumerate(block.data):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(7)
                if index == 0:
                    p.paragraph_format.keep_with_next = len(block.data) > 1
                if is_content_lead(text):
                    p.paragraph_format.keep_with_next = True
                # Bold lead-in paragraphs replace former shaded cards.
                add_inline(p, text, BODY_SIZE, INK, index == 0)
        elif block.kind in ("bullet", "number"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.17 + block.level * .17)
            p.paragraph_format.first_line_indent = Inches(-.13)
            p.paragraph_format.space_after = Pt(5)
            prefix = "• " if block.kind == "bullet" else f"{block.data}. "
            item = block.text
            checked = re.match(r"^\[([ xX])\]\s*(.*)$", item)
            if checked:
                # Editable text check targets remain readable in NanumSquare.
                prefix = "[x] " if checked.group(1).lower() == "x" else "[ ] "
                item = checked.group(2)
            add_inline(p, prefix + item)
            if is_content_lead(item):
                p.paragraph_format.keep_with_next = True
        else:
            p = doc.add_paragraph()
            add_inline(p, block.text)
            if (is_content_lead(block.text)
                    or block.text == "아래 항목으로 메뉴의 역할과 연결 관계를 이해했는지 점검합니다."
                    or (in_menu_chapter and identifier < len(blocks)
                        and blocks[identifier].kind == "table")):
                p.paragraph_format.keep_with_next = True
    # A short completion checklist is one reading unit. Keep its heading and
    # items together rather than leaving one or two checks on a trailing page.
    for index, paragraph in enumerate(doc.paragraphs[:-1]):
        if paragraph.text.startswith(("[ ] ", "[x] ")):
            next_paragraph = doc.paragraphs[index + 1]
            paragraph.paragraph_format.keep_together = True
            if next_paragraph.text.startswith(("[ ] ", "[x] ")):
                paragraph.paragraph_format.keep_with_next = True
    enforce_font(doc)
    doc.core_properties.title = "Mold 신입 개발자 교육 안내서"
    doc.core_properties.subject = "macOS 개발 환경과 Mold 개발 및 검증 절차"
    doc.core_properties.author = "ABLESTACK Mold 개발팀"
    doc.core_properties.keywords = "Mold, ABLESTACK, 온보딩, macOS"
    doc.core_properties.comments = ""
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    embed_document_fonts(docx_path)
    print(f"[OK] Generated {docx_path}")
    print(f"[OK] Font {FONT_NAME}; {len(entries)} TOC links; "
          f"{sum(block.kind == 'table' for block in blocks)} comparison tables")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path,
                        default=HERE / "mold-developer-onboarding-guide.md")
    parser.add_argument("--output", type=Path,
                        default=HERE / "Mold 신입 개발자 교육 안내서 v1.0.docx")
    parser.add_argument("--page-map", type=Path,
                        default=HERE / "toc-page-map.json",
                        help="Verified bookmark-to-page cache from the latest render")
    args = parser.parse_args()
    page_map = json.loads(args.page_map.read_text()) if args.page_map.is_file() else None
    create_styled_document(args.input, args.output, page_map)


if __name__ == "__main__":
    main()
